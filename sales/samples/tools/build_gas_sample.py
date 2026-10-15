"""GASサンプルの表(ダミーデータ)を作る。同じデータでPython側の期待値JSONも出力する。"""
import json, random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

random.seed(7)
ITEMS = [("ハンドメイド雑貨A", 1200), ("ハンドメイド雑貨B", 2500), ("ギフトセット", 4800)]
NAMES = ["サンプル太郎", "サンプル花子", "テスト一郎", "テスト二葉", "ダミー三郎"]
rows = []
for i in range(36):
    month = [8, 9, 10][i % 3]
    day = random.randint(1, 28)
    t = dt.datetime(2026, month, day, random.randint(9, 20), random.randint(0, 59))
    item, price = random.choice(ITEMS)
    rows.append([t, random.choice(NAMES), item, random.randint(1, 4), price])
rows.sort(key=lambda r: r[0])

# 不正行(テスト用): 数量が小数・日時なし・商品名なし
bad = [[None, "ダミー", "ギフトセット", 1, 4800], [dt.datetime(2026, 9, 5, 10, 0), "ダミー", "", 2, 1200],
       [dt.datetime(2026, 9, 6, 10, 0), "ダミー", "ギフトセット", 1.5, 4800]]

def expected(rs):
    tot = {}
    for t, n, item, q, p in rs:
        if t is None or not item or int(q) != q or q <= 0: continue
        k = (t.strftime("%Y-%m"), item)
        a = tot.setdefault(k, [0, 0]); a[0] += q; a[1] += q * p
    return [dict(month=k[0], item=k[1], qty=v[0], amount=v[1]) for k, v in sorted(tot.items())]

json.dump(dict(
    rows=[[r[0].isoformat() if r[0] else None, r[1], r[2], r[3], r[4]] for r in rows + bad],
    expected=expected(rows + bad), skipped=len(bad)), open("../gas/fixture.json", "w"), ensure_ascii=False, indent=1)

B = Font(bold=True); GRY = PatternFill("solid", fgColor="D9D9D9"); Y = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="999999"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()
ws = wb.active; ws.title = "フォーム回答"
for i, h in enumerate(["日時", "氏名", "商品", "数量", "単価", "金額", "年月"], 1):
    c = ws.cell(1, i, h); c.font = B; c.fill = GRY; c.border = BOX; c.alignment = Alignment(horizontal="center")
for r, row in enumerate(rows, 2):
    for i, v in enumerate(row, 1):
        c = ws.cell(r, i, v); c.border = BOX
    ws.cell(r, 1).number_format = "yyyy/mm/dd hh:mm"; ws.cell(r, 5).number_format = '#,##0'
    ws.cell(r, 6, f"=D{r}*E{r}").number_format = '#,##0'
    ws.cell(r, 7, f'=TEXT(A{r},"yyyy-mm")')
    for col in range(6, 8): ws.cell(r, col).border = BOX
last = len(rows) + 1
for col, w in zip("ABCDEFG", [18, 14, 18, 8, 10, 12, 10]): ws.column_dimensions[col].width = w
ws["I1"] = "※ダミーデータ。実際はGoogleフォームの回答がA〜E列に入ります(F・G列は表示用の計算列)。"

sm = wb.create_sheet("月次集計")
sm["A1"] = "月次集計(GAS実行後の出力イメージ)"; sm["A1"].font = Font(bold=True, size=14)
hdr = ["年月", "商品", "数量", "金額(円)"]
for i, h in enumerate(hdr, 1):
    c = sm.cell(3, i, h); c.font = B; c.fill = GRY; c.border = BOX; c.alignment = Alignment(horizontal="center")
r = 4
months = ["2026-08", "2026-09", "2026-10"]
for m in months:
    for item, _ in ITEMS:
        sm.cell(r, 1, m); sm.cell(r, 2, item)
        sm.cell(r, 3, f'=SUMIFS(フォーム回答!$D$2:$D${last},フォーム回答!$G$2:$G${last},A{r},フォーム回答!$C$2:$C${last},B{r})')
        sm.cell(r, 4, f'=SUMIFS(フォーム回答!$F$2:$F${last},フォーム回答!$G$2:$G${last},A{r},フォーム回答!$C$2:$C${last},B{r})').number_format = '#,##0'
        for col in range(1, 5): sm.cell(r, col).border = BOX
        r += 1
sm.cell(r, 2, "合計").font = B
sm.cell(r, 3, f"=SUM(C4:C{r-1})").font = B
c = sm.cell(r, 4, f"=SUM(D4:D{r-1})"); c.font = B; c.number_format = '#,##0'
for col in range(1, 5): sm.cell(r, col).border = BOX
sm.cell(r + 2, 1, "GASを実行すると、この表が自動で更新され、月初にメールで届きます。")
for col, w in zip("ABCD", [12, 22, 8, 14]): sm.column_dimensions[col].width = w
sm.page_setup.orientation = "portrait"; sm.page_setup.fitToWidth = 1; sm.sheet_properties.pageSetUpPr.fitToPage = True
wb.save("../gas/sample-form-summary.xlsx")
print("rows", len(rows), "summary last row", r)
