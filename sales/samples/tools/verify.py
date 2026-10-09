# 検証: 入力を差し替えたコピーをLibreOfficeで再計算し、Python整数モデルと比較
import os, subprocess, sys, tempfile, datetime, random
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__)); S = os.path.dirname(HERE)
tmp = tempfile.mkdtemp()
def model(items, mode, wh):
    sub = {10: 0, 8: 0}
    for q, p, r in items: sub[r] += int(q * p)
    k = {"切捨て": 0, "四捨五入": 50, "切上げ": 99}[mode]
    t = {r: (sub[r] * r + k) // 100 for r in sub}
    base = sub[10] + sub[8]
    w = (1021 * min(base, 10**6) + 2042 * max(base - 10**6, 0)) // 10000 if wh == "する" else 0
    return dict(F33=sub[10], F34=t[10], F35=sub[8], F36=t[8], F37=base, F38=t[10]+t[8], F39=base+t[10]+t[8], F40=w, F41=base+t[10]+t[8]-w)
cases = []
def case(items, mode="切捨て", wh="しない"): cases.append((items, mode, wh))
case([(1, 100000, 10)])
for m in ("切捨て", "四捨五入", "切上げ"): case([(1, 333, 8)], m)
case([(1, 1200000, 10)], "切捨て", "する")
case([(1, 1000000, 10)], "切捨て", "する")
case([(1, 1000001, 10)], "切捨て", "する")
case([(1, 12500, 8)], "切上げ")   # 浮動小数点の罠(12500*0.08)
case([(3, 111, 10), (7, 49, 8), (2, 1005, 10), (1, 1, 8)], "四捨五入", "する")
random.seed(1)
for _ in range(15):
    n = random.randint(1, 8)
    case([(random.randint(1, 9), random.randint(1, 300000), random.choice((10, 8))) for _ in range(n)], random.choice(("切捨て","四捨五入","切上げ")), random.choice(("する","しない")))
files = []
for kind in ("invoice-template.xlsx", "quote-template.xlsx"):
    for i, (items, mode, wh) in enumerate(cases):
        wb = load_workbook(os.path.join(S, kind)); ws = wb.worksheets[0]
        for r in range(12, 32):
            for c in range(1, 5): ws.cell(r, c).value = None
        for j, (q, p, rt) in enumerate(items):
            r = 12 + j; ws.cell(r, 1).value = f"品{j}"; ws.cell(r, 2).value = q; ws.cell(r, 3).value = p; ws.cell(r, 4).value = rt
        ws["F8"] = mode; ws["F9"] = wh
        f = os.path.join(tmp, "in", f"{kind[:3]}_{i:02d}.xlsx"); os.makedirs(os.path.dirname(f), exist_ok=True); wb.save(f); files.append((kind, i, f))
# tracker
wb = load_workbook(os.path.join(S, "expense-tracker.xlsx")); ws = wb["取引"]
tx = [(datetime.date(2026,1,5),"売上","",100000,None),(datetime.date(2026,1,6),"経費","通信費",5555,50),
      (datetime.date(2026,2,6),"経費","通信費",3333,33),(datetime.date(2026,2,7),"経費","消耗品費",1000,None),
      (datetime.date(2026,12,31),"売上","",50000,None),(datetime.date(2025,12,31),"売上","",99999,None),(datetime.date(2026,3,1),"経費","地代家賃",80000,30)]
for r in range(5, 12): 
    for c in range(1, 7): ws.cell(r, c).value = None
for j, (d, k, a, amt, p) in enumerate(tx):
    r = 5 + j; ws.cell(r,1).value=d; ws.cell(r,2).value=k; ws.cell(r,3).value=a or None; ws.cell(r,5).value=amt; ws.cell(r,6).value=p
ft = os.path.join(tmp, "in", "trk.xlsx"); wb.save(ft)
out = os.path.join(tmp, "out")
subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, *[f for _,_,f in files], ft], check=True, capture_output=True)
bad = 0
for kind, i, f in files:
    v = load_workbook(os.path.join(out, os.path.basename(f)), data_only=True).worksheets[0]
    exp = model(*cases[i])
    for cell, e in exp.items():
        if v[cell].value != e: bad += 1; print("NG", kind, i, cases[i], cell, v[cell].value, e)
print("請求書/見積書 cases:", len(files), "mismatches:", bad)
v = load_workbook(os.path.join(out, "trk.xlsx"), data_only=True)["年間集計"]
exp_m = {1:(100000, 2777+0),2:(0,1099+1000),3:(0,24000),12:(50000,0)}
# 5555*50%=2777, 3333*33%=1099(1099.89切捨て), 80000*30%=24000
tb = 0
for m in range(1, 13):
    e = exp_m.get(m, (0, 0))
    if (v.cell(4+m,2).value, v.cell(4+m,3).value) != e: tb += 1; print("NG月", m, v.cell(4+m,2).value, v.cell(4+m,3).value, e)
for cell, e in {"B17":150000, "C17":2777+2099+24000}.items():
    if v[cell].value != e: tb += 1; print("NG", cell, v[cell].value, e)
print("通信費", v["B21"].value, v["C21"].value, v["D21"].value, "(期待 3876/8888/5012)", v["B34"].value)
if (v["B21"].value, v["C21"].value, v["D21"].value) != (3876, 8888, 5012): tb += 1
print("tracker mismatches:", tb)
sys.exit(1 if bad or tb else 0)
