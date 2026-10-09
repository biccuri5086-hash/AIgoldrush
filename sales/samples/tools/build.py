# サンプル生成スクリプト: python3 build.py  (出力は一つ上のディレクトリ)
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
YEL = PatternFill("solid", fgColor="FFF2CC")
GRY = PatternFill("solid", fgColor="D9D9D9")
B = Font(bold=True)
thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
YEN = '#,##0'

def inp(c, fmt=None):
    c.fill = YEL; c.protection = Protection(locked=False); c.border = BOX
    if fmt: c.number_format = fmt

def lab(c):
    c.font = B

def doc_sheet(wb, title, lines):
    ws = wb.create_sheet(title)
    ws.column_dimensions["A"].width = 110
    for i, t in enumerate(lines, 1):
        ws.cell(i, 1, t).alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1: ws.cell(i, 1).font = Font(bold=True, size=14)
    return ws

def doc(kind):
    q = kind == "見積書"
    return [
        f"{kind}テンプレート 使い方",
        "1. 黄色のセルだけが入力欄です。それ以外(数式セル)はシート保護(パスワードなし)でロックしています。構成を変えるときは「校閲 > シート保護の解除」。",
        "2. 明細は最大20行。税率は 10 または 8 をリストから選びます(8は軽減税率。飲食料品・定期購読の新聞など)。",
        "3. 金額欄 = 数量 x 単価(1円未満は切捨て)。消費税額は『税率ごとに合計した税抜金額』に対して1回だけ計算します(明細行ごとには計算しません)。",
        "4. 端数処理(切捨て/四捨五入/切上げ)は税率ごとの消費税額に適用されます。整数演算で計算するため浮動小数点誤差は出ません。",
        "5. 登録番号は T + 数字13桁。入力欄の隣に形式チェックを出しますが、実在確認は国税庁の適格請求書発行事業者公表サイトで行ってください。",
        "6. 源泉徴収は『する/しない』のスイッチ。対象は税抜金額: 100万円以下の部分 10.21%、100万円超の部分 20.42%、1円未満切捨て。"
        "報酬と消費税が区分表示されている場合に税抜で計算できる取扱いです(国税庁 源泉所得税の質疑応答事例等)。対象報酬の種類かどうかは必ず国税庁の最新情報で確認してください。",
        "7. 印刷はA4縦・横1ページ幅に設定済みです。",
        "8. サンプルの氏名・登録番号・口座などはすべてダミーです。実在しません。",
        "限界: マイナス行(値引)を含む場合は税率ごとの小計がマイナスにならないよう注意。20行超、複数ページ、インボイスの電子保存要件、個別事情の税務判断は対象外です。税務上の最終判断は税理士等へ。"
        + ("" if not q else " 見積書の源泉徴収額は参考表示です。"),
    ]

def build_doc(kind, fname):
    q = kind == "見積書"
    wb = Workbook(); ws = wb.active; ws.title = kind
    for col, w in zip("ABCDEF", (36, 9, 13, 8, 6, 15)):
        ws.column_dimensions[col].width = w
    ws["A1"] = kind; ws["A1"].font = Font(bold=True, size=18)
    ws.merge_cells("A1:F1"); ws["A1"].alignment = Alignment(horizontal="center")
    meta = [
        (3, "宛名(御中)", "B3", "株式会社サンプル", "E3", "番号", "F3", ("Q" if q else "INV") + "-0001"),
        (4, "件名" if q else "請求元(屋号/氏名)", "B4", "ホームページ制作" if q else "サンプル商店(ダミー)", "E4", "発行日", "F4", None),
        (5, "請求元住所" if not q else "発行者(屋号/氏名)", "B5", "東京都千代田区サンプル1-1-1(ダミー)" if not q else "サンプル商店(ダミー)", "E5", "登録番号", "F5", "T0000000000000"),
        (8, "有効期限" if q else "お支払期限", "B8", None, "E8", "端数処理", "F8", "切捨て"),
        (9, "納期・備考" if q else "振込先", "B9", "" if q else "サンプル銀行 本店 普通 0000000(ダミー)", "E9", "源泉徴収", "F9", "しない"),
    ]
    for r, l1, c1, v1, l2n, l2, c2, v2 in meta:
        ws.cell(r, 1, l1); lab(ws.cell(r, 1))
        ws[c1] = v1; ws.merge_cells(f"B{r}:D{r}"); inp(ws[c1])
        ws[l2n] = l2; lab(ws[l2n]); ws[c2] = v2; inp(ws[c2])
    ws["F4"].number_format = "yyyy/mm/dd"
    if q: ws["B8"].number_format = "yyyy/mm/dd"
    else: ws["B8"].number_format = "yyyy/mm/dd"
    ws["A6"] = "登録番号チェック"; lab(ws["A6"])
    ws["B6"] = ('=IF(F5="","未入力",IF(IFERROR(AND(LEN(F5)=14,LEFT(F5,1)="T",VALUE(MID(F5,2,13))>=0,ISNUMBER(VALUE(MID(F5,2,13)))),FALSE),'
                '"形式OK(実在は国税庁サイトで確認)","形式エラー: Tと数字13桁"))')
    ws.merge_cells("B6:D6")
    ws["A10"] = "※は軽減税率(8%)対象"; 
    for i, h in enumerate(["品名","数量","単価","税率(%)","軽減","金額(税抜)"], 1):
        c = ws.cell(11, i, h); c.font = B; c.fill = GRY; c.border = BOX; c.alignment = Alignment(horizontal="center")
    F, L = 12, 31
    sample = {12: ("Webサイト制作", 1, 100000, 10), 13: ("飲食料品(サンプル)", 1, 333, 8)}
    for r in range(F, L + 1):
        s = sample.get(r)
        for col in range(1, 5):
            c = ws.cell(r, col, s[col - 1] if s else None); inp(c)
        ws.cell(r, 2).number_format = "#,##0.##"; ws.cell(r, 3).number_format = YEN
        ws.cell(r, 5, f'=IF(D{r}=8,"※","")').alignment = Alignment(horizontal="center")
        ws.cell(r, 6, f'=IF(OR(B{r}="",C{r}=""),0,SIGN(B{r}*C{r})*ROUNDDOWN(ROUND(ABS(B{r}*C{r}),6),0))').number_format = YEN
        for col in (5, 6): ws.cell(r, col).border = BOX
    dv = DataValidation(type="list", formula1='"10,8"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"D{F}:D{L}")
    dv2 = DataValidation(type="list", formula1='"切捨て,四捨五入,切上げ"', allow_blank=False); ws.add_data_validation(dv2); dv2.add("F8")
    dv3 = DataValidation(type="list", formula1='"する,しない"', allow_blank=False); ws.add_data_validation(dv3); dv3.add("F9")
    ws["A32"] = f'=IF(SUMPRODUCT((F{F}:F{L}<>0)*(D{F}:D{L}<>10)*(D{F}:D{L}<>8))>0,"警告: 税率未選択の行があります","")'
    ws["A32"].font = Font(bold=True, color="C00000")
    k = '(IF($F$8="切捨て",0,IF($F$8="四捨五入",50,99)))'
    def tax(sub, r): return f'=SIGN({sub})*INT((ABS({sub})*{r}+{k})/100)'
    rows = [
        (33, "10%対象 税抜合計", f"=SUMIF(D{F}:D{L},10,F{F}:F{L})"),
        (34, "10%対象 消費税額", tax("F33", 10)),
        (35, "8%対象(※) 税抜合計", f"=SUMIF(D{F}:D{L},8,F{F}:F{L})"),
        (36, "8%対象(※) 消費税額", tax("F35", 8)),
        (37, "税抜合計", "=F33+F35"),
        (38, "消費税合計", "=F34+F36"),
        (39, "税込合計", "=F37+F38"),
        (40, "源泉徴収税額" + ("(参考)" if q else ""),
         '=IF($F$9="する",SIGN(F37)*INT((1021*MIN(ABS(F37),1000000)+2042*MAX(ABS(F37)-1000000,0))/10000),0)'),
        (41, "見積金額(税込)" if q else "ご請求額(税込−源泉徴収)", "=F39-F40"),
    ]
    for r, l, f in rows:
        ws.cell(r, 1, l).font = B; c = ws.cell(r, 6, f); c.number_format = YEN; c.border = BOX
        ws.merge_cells(f"A{r}:E{r}"); ws.cell(r, 1).alignment = Alignment(horizontal="right")
    ws["F41"].font = Font(bold=True, size=12)
    ws["A43"] = "備考"; lab(ws["A43"]); ws["B43"] = "(自由記入)"; ws.merge_cells("B43:F45"); inp(ws["B43"])
    ws["B43"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.print_area = "A1:F45"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4; ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.protection.sheet = True
    doc_sheet(wb, "使い方", doc(kind))
    wb.save(os.path.join(OUT, fname))

def build_tracker():
    wb = Workbook(); ws = wb.active; ws.title = "取引"
    ys = wb.create_sheet("年間集計")
    ys["A1"] = "年間集計"; ys["A1"].font = Font(bold=True, size=14)
    ys["A2"] = "対象年(西暦)"; lab(ys["A2"]); ys["B2"] = 2026; inp(ys["B2"], "0")
    ys.column_dimensions["A"].width = 18
    for c in "BCDE": ys.column_dimensions[c].width = 16
    for i, h in enumerate(["月", "売上", "経費(按分後)", "差引"], 1):
        c = ys.cell(4, i, h); c.font = B; c.fill = GRY; c.border = BOX
    N0, N1 = 5, 304
    R = lambda col: f"取引!${col}${N0}:${col}${N1}"
    for m in range(1, 13):
        r = 4 + m
        ys.cell(r, 1, m)
        ys.cell(r, 2, f'=SUMIFS({R("G")},{R("B")},"売上",{R("H")},A{r})')
        ys.cell(r, 3, f'=SUMIFS({R("G")},{R("B")},"経費",{R("H")},A{r})')
        ys.cell(r, 4, f"=B{r}-C{r}")
    ys["A17"] = "年間合計"; lab(ys["A17"])
    for col in "BCD": ys[f"{col}17"] = f"=SUM({col}5:{col}16)"
    for r in range(5, 18):
        for col in "BCD": ys[f"{col}{r}"].number_format = YEN; ys[f"{col}{r}"].border = BOX
    for i, h in enumerate(["勘定科目(編集可)", "経費計上額", "支出総額", "家事按分で除外した額"], 1):
        c = ys.cell(20, i, h); c.font = B; c.fill = GRY; c.border = BOX
    ys.column_dimensions["A"].width = 24
    accts = ["通信費","水道光熱費","地代家賃","消耗品費","旅費交通費","接待交際費","外注工賃","広告宣伝費","租税公課","減価償却費","雑費","その他"]
    for i, a in enumerate(accts):
        r = 21 + i
        ys.cell(r, 1, a); inp(ys.cell(r, 1))
        ys.cell(r, 2, f'=SUMIFS({R("G")},{R("B")},"経費",{R("C")},A{r},{R("H")},">0")')
        ys.cell(r, 3, f'=SUMIFS({R("E")},{R("B")},"経費",{R("C")},A{r},{R("H")},">0")')
        ys.cell(r, 4, f"=C{r}-B{r}")
        for col in "BCD": ys[f"{col}{r}"].number_format = YEN; ys[f"{col}{r}"].border = BOX
    ys["A33"] = "科目合計"; lab(ys["A33"])
    for col in "BCD": ys[f"{col}33"] = f"=SUM({col}21:{col}32)"; ys[f"{col}33"].number_format = YEN
    ys["A34"] = "チェック"; ys["B34"] = '=IF(B33=C17,"OK: 月次経費と一致","不一致: 科目未設定の経費あり")'
    ys.protection.sheet = True
    ys.page_setup.paperSize = ys.PAPERSIZE_A4; ys.page_setup.fitToWidth = 1; ys.page_setup.fitToHeight = 1
    ys.sheet_properties.pageSetUpPr.fitToPage = True
    # 取引
    hs = ["日付", "区分", "勘定科目", "摘要", "金額(税込)", "事業割合(%)", "計上額", "月"]
    for i, h in enumerate(hs, 1):
        c = ws.cell(4, i, h); c.font = B; c.fill = GRY; c.border = BOX
    ws["A1"] = "売上・経費 取引入力"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "黄色欄に入力。事業割合は経費のみ有効(空欄=100)。家事按分後の額は1円未満切捨て。"
    for col, w in zip("ABCDEFGH", (12, 8, 14, 30, 14, 12, 14, 6)): ws.column_dimensions[col].width = w
    import datetime
    sample = [(datetime.date(2026,1,10),"売上","","サンプル案件A(ダミー)",110000,None),
              (datetime.date(2026,1,15),"経費","通信費","インターネット回線(ダミー)",5500,50),
              (datetime.date(2026,2,3),"経費","消耗品費","文具(ダミー)",1234,None)]
    for r in range(N0, N1 + 1):
        s = sample[r - N0] if r - N0 < len(sample) else (None,) * 6
        for col, v in enumerate(s, 1):
            c = ws.cell(r, col, v); inp(c)
        ws.cell(r, 1).number_format = "yyyy/mm/dd"; ws.cell(r, 5).number_format = YEN
        ws.cell(r, 7, f'=IF(OR(E{r}="",B{r}=""),0,IF(B{r}="売上",E{r},SIGN(E{r})*INT(ABS(E{r})*IF(F{r}="",100,F{r})/100)))').number_format = YEN
        ws.cell(r, 8, f'=IF(A{r}="","",IF(YEAR(A{r})=年間集計!$B$2,MONTH(A{r}),0))')
    d1 = DataValidation(type="list", formula1='"売上,経費"', allow_blank=True); ws.add_data_validation(d1); d1.add(f"B{N0}:B{N1}")
    d2 = DataValidation(type="list", formula1="=年間集計!$A$21:$A$32", allow_blank=True); ws.add_data_validation(d2); d2.add(f"C{N0}:C{N1}")
    d3 = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True); ws.add_data_validation(d3); d3.add(f"F{N0}:F{N1}")
    ws.freeze_panes = "A5"; ws.protection.sheet = True
    ws.print_area = f"A1:H{N1}"; ws.page_setup.paperSize = ws.PAPERSIZE_A4; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    doc_sheet(wb, "使い方", [
        "売上・経費管理 使い方",
        "1. 「年間集計」B2に対象年を入れ、「取引」に1行ずつ入力します(最大300行)。対象年以外の日付の行は集計されません。",
        "2. 区分=経費の行で『事業割合(%)』を入れると、金額 x 割合 を1円未満切捨てして計上額にします。空欄は100%です。",
        "3. 勘定科目は「年間集計」A21:A32のリストから選びます。リスト名は編集可能です。",
        "4. 金額は税込で入力する前提です。消費税の区分・簡易課税・インボイスの仕入税額控除の集計はしません。",
        "5. 家事按分の割合の妥当性(使用実態に基づく合理的な根拠)は利用者の判断です。税務判断は税理士等へ。",
        "6. 氏名・金額はすべてダミーです。",
    ])
    wb.save(os.path.join(OUT, "expense-tracker.xlsx"))

build_doc("請求書", "invoice-template.xlsx")
build_doc("見積書", "quote-template.xlsx")
build_tracker()
